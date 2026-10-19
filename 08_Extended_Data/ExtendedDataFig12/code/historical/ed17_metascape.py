"""Import, compose and publish the supplied nine-list Metascape ED17 results.

Run with /usr/bin/python3; --import is needed only for the initial import.
The packaged data are sufficient for --render and --publish.
"""
from pathlib import Path
from copy import deepcopy
import base64,csv,hashlib,json,re,shutil,sys
from datetime import datetime,timezone
from urllib.parse import unquote
B=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(B/'Code/vendor'))
import openpyxl
import pymupdf as fitz
from lxml import etree,html
from extended_layout import align_page

D=B/'Source_data/metascape_20261010'
ART=B/'Assets/arial'
PARSER=html.HTMLParser(huge_tree=True)
FONT='/usr/share/fonts/truetype/msttcorefonts/Arial.ttf'
BOLD='/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf'
METHODS=['DeepVelo(GCN)','LatentVelo','STT','UniTVelo','biVI','dynamo','scVelo(dyn)','scVelo(sto)','veloVI']
RUNS={100:'all.tsuluz1m0',200:'all.to4our0de'}
STEM='Extended_Data_Fig17'
TITLE='Top100 and top200 gene overlap and enrichment'
COMMON_EN='Metascape used human Entrez mapping (database 2026-07-01), a whole-genome background, P < 0.01, at least three hits and enrichment factor > 1.5; q-values use Benjamini-Hochberg correction. Purple chords link identical mapped genes; dark/light orange denotes shared/list-specific genes. Heatmaps retain each report’s 20 representative clusters, independent clustering and -log10(P) scale; grey denotes zero/no enrichment in the supplied matrix. List definitions differ across methods. biVI ranks Radial Glia-versus-other-cell norm_spliced_mean differential results by Bayes factor, including both LFC signs; this is not a direct velocity-gene ranking. These are descriptive gene-list comparisons, not method-accuracy tests.'
COMMON_ZH="Metascape used human Entrez mapping (database 2026-07-01), a whole-genome background, P < 0.01, at least three hits and enrichment factor > 1.5; q-values use Benjamini-Hochberg correction. Purple chords link identical mapped genes; dark/light orange denotes shared/list-specific genes. Heatmaps retain each report's 20 representative clusters, independent clustering and -log10(P) scale; grey denotes zero/no enrichment in the supplied matrix. List definitions differ across methods. biVI ranks Radial Glia-versus-other-cell norm_spliced_mean differential results by Bayes factor, including both LFC signs; this is not a direct velocity-gene ranking. These are descriptive gene-list comparisons, not method-accuracy tests. "
CAPTION=('Top row, Gene overlap and enrichment for nine-configuration top100 lists (99–100 mapped genes per list; 519 distinct mapped genes). Middle row, Corresponding top200 results (198–200 mapped genes per list; 809 distinct mapped genes). '+COMMON_EN+' Bottom row, Historical young and aged ReDeeM LatentVelo fields, respectively, using the reported 784 shared genes; these are separate from the forebrain enrichment analyses. The exact 784-gene manifest and matching run remain unresolved. No new aging comparison or velocity model fitting was performed.\n'+'Top row, Gene overlap and enrichment for nine-configuration top100 lists (99-100 mapped genes per list; 519 distinct mapped genes). Middle row, Corresponding top200 results (198-200 mapped genes per list; 809 distinct mapped genes). '+COMMON_ZH+'Bottom left and right, Historical young and aged ReDeeM LatentVelo fields using the reported 784 shared genes, separate from the forebrain enrichment analyses. The exact gene list and matching run remain unresolved. No new aging comparison or velocity model fitting was performed.')

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def jwrite(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n')
def readcsv(p):
    with p.open() as f:return list(csv.DictReader(f))
def csvwrite(p,rows):
    with p.open('w',newline='') as f:csv.writer(f).writerows(rows)
def parse(p):return html.document_fromstring(p.read_bytes(),parser=PARSER)
def write(p,d):p.write_bytes(html.tostring(d,encoding='utf-8',doctype='<!doctype html>'))
def savepdf(doc,p):
    tmp=p.with_suffix('.tmp.pdf');doc.save(tmp,garbage=4,deflate=True);doc.close();tmp.replace(p)

def import_data():
    D.mkdir(parents=True,exist_ok=True)
    report={'date':'2026-10-10','database_date':'2026-07-01','background':'all genes in the genome','organism':'H. sapiens','p_cutoff':0.01,'min_hits':3,'min_enrichment':1.5,'q_value_correction':'Benjamini-Hochberg','term_cluster_similarity':0.3,'runs':{}}
    for n,run in RUNS.items():
        src=B.parent/'dynamic_gene_metascape'/run;dst=D/f'top{n}';dst.mkdir(exist_ok=True)
        paths=['AnalysisReport.html','metascape_result.xlsx','README.txt']
        for folder in ['Overlap_circos','Enrichment_heatmap','Enrichment_GO']:
            paths.extend(str(p.relative_to(src)) for p in (src/folder).iterdir() if p.suffix in ['.csv','.cdt','.atr','.gtr','.pdf','.svg','.png'])
        hashes={}
        for rel in paths:
            target=dst/rel;target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src/rel,target);hashes[rel]=sha(target)
        inputs=B.parent/'metascape_inputs_20261010_with_biVI'/f'Metascape_Forebrain_top{n}_9methods.csv'
        shutil.copy2(inputs,dst/'submitted_gene_lists.csv')
        incoming=readcsv(inputs)
        w=openpyxl.load_workbook(dst/'metascape_result.xlsx',read_only=True,data_only=True)
        rows=list(w['Annotation'].values);header=rows[0];rows=rows[1:];w.close()
        # Preserve server mapping, including merged aliases, without substituting another annotation.
        mapped=[['Entrez_ID','original_id','Gene Symbol']+METHODS]
        for r in rows:mapped.append([r[0],r[1],r[header.index('Gene Symbol')]]+[int(r[header.index(re.sub(r'[^\w]','_',m))]) for m in METHODS])
        csvwrite(dst/'mapped_membership.csv',mapped)
        counts=[];audit=[]
        for method in METHODS:
            exp={r[method] for r in incoming};col=header.index(re.sub(r'[^\w]','_',method));seen=set();ids=set()
            for r in rows:
                aliases=set(str(r[1]).split(';'));expected=bool(exp&aliases);actual=int(r[col])==1
                assert expected==actual,(n,method,r[:3])
                if actual:seen.update(exp&aliases);ids.add(r[0])
            missing=sorted(exp-seen)
            assert len(exp)==n
            counts.append([method,n,len(ids),';'.join(missing),len(seen)-len(ids)])
            for gene in missing:audit.append([method,gene,'not represented in returned annotation'])
        csvwrite(dst/'list_counts.csv',[['method','submitted','mapped_unique','unmapped_symbols','merged_identifier_reduction']]+counts)
        csvwrite(dst/'unmapped_symbols.csv',[['method','gene','status']]+audit)
        fulltext=parse(dst/'AnalysisReport.html').text_content()
        assert '2026-07-01' in fulltext and 'All genes in the genome' in fulltext
        report['runs'][str(n)]={'archive':str((src.parent/(run+'.zip')).relative_to(B.parent)),'archive_sha256':sha(src.parent/(run+'.zip')),'mapped_union':len(rows),'submitted_union':len({v for r in incoming for v in r.values()}),'lists':[dict(zip(['method','submitted','mapped_unique','unmapped_symbols','merged_identifier_reduction'],r)) for r in counts],'copied_source_sha256':hashes,'input_sha256':sha(inputs),'mapping_membership_verified':True}
    shutil.copy2(B.parent/'metascape_inputs_20261010_with_biVI/biVI_RadialGlia_ranking_audit.csv',D/'biVI_RadialGlia_ranking_audit.csv')
    # Freeze the exact pre-update source for the unchanged historical age-group panels.
    hist=B.parent/'Manuscript_1010_history/before_ed17_metascape_20261010/Assets/arial/Extended_Data_Fig17.pdf'
    shutil.copy2(hist,D/'historical_ED17.pdf')
    report['historical_age_source_sha256']=sha(D/'historical_ED17.pdf')
    jwrite(D/'import_manifest.json',report)
    jwrite(B/'Provenance/ed17_metascape.json',report)
    print('Imported and verified nine-list mappings for top100 and top200.',flush=True)

def text(page,x,y,s,size=8,bold=False):
    page.insert_text((x,y),s,fontname='ED17ArialBold' if bold else 'ED17Arial',fontfile=BOLD if bold else FONT,fontsize=size)

def centered_text(page,cx,y,s,size=8,bold=False):
    font=fitz.Font(fontfile=BOLD if bold else FONT)
    text(page,cx-font.text_length(s,fontsize=size)/2,y,s,size,bold)

def circle_pdf(path):
    tree=etree.fromstring(path.read_bytes(),parser=etree.XMLParser(resolve_entities=False,no_network=True))
    labels=tree.xpath('//*[local-name()="text"]')
    specs=[]
    for t in labels:
        specs.append((float(t.get('x')),float(t.get('y')),float(t.get('font-size').replace('px','')),float(re.search(r'rotate\(([-\d.]+)',t.get('transform'))[1]),''.join(t.itertext())))
        t.getparent().remove(t)
    svg=fitz.open(stream=etree.tostring(tree),filetype='svg');doc=fitz.open('pdf',svg.convert_to_pdf());svg.close()
    font=fitz.Font(fontfile=FONT)
    for x,y,size,angle,label in specs:
        doc[0].insert_text((x-font.text_length(label,fontsize=size)/2,y),label,fontname='CircleArial',fontfile=FONT,fontsize=size,morph=(fitz.Point(x,y),fitz.Matrix(-angle)))
    return doc

def render():
    assert (D/'import_manifest.json').exists()
    out=B/'Reports/ed17_metascape';out.mkdir(exist_ok=True)
    hist=fitz.open(D/'historical_ED17.pdf')
    # Remove only old d/e letters; all point and streamline geometry stays intact.
    for block in hist[0].get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            for span in line['spans']:
                if (span['text'].strip() in ['d','e'] and span['bbox'][1]>510) or ('DeepVelo' in span['text'] and span['bbox'][3]>512):
                    hist[0].add_redact_annot(fitz.Rect(span['bbox']),fill=(1,1,1))
    hist[0].apply_redactions(images=0,graphics=0)
    doc=fitz.open();page=doc.new_page(width=595.276,height=841.89)
    centered_text(page,page.rect.width/2,20,'EXTENDED DATA FIGURE 17',8,True)
    for n,y in [(100,48),(200,324)]:
        centered_text(page,page.rect.width/2,y,f'Top{n}',10,True)
        centered_text(page,126,y+22,'Gene overlap',8)
        centered_text(page,402.5,y+22,'Enrichment',8)
        src=circle_pdf(D/f'top{n}/Overlap_circos/CircosOverlapByGene.svg')
        page.show_pdf_page(fitz.Rect(22,y+49,230,y+257),src,0);src.close()
        with fitz.open(D/f'top{n}/Enrichment_heatmap/HeatmapSelectedGO.pdf') as src:
            height=343*src[0].rect.height/src[0].rect.width
            page.show_pdf_page(fitz.Rect(231,y+34,574,y+34+height),src,0)
    centered_text(page,page.rect.width/2,606,'Historical age-group fields | LatentVelo',10,True)
    # Keep the complete original fields, Young/aged labels and cell-type legend.
    page.show_pdf_page(fitz.Rect(58,617,557,779.45),hist,0,clip=fitz.Rect(80,512,507,651))
    text(page,24,803,'Purple chords: identical genes. Orange inner ring: shared / list-specific genes.',7)
    doc.subset_fonts()
    savepdf(doc,ART/(STEM+'.pdf'))
    alignment=align_page(ART/(STEM+'.pdf'),B/'Extended_Data/figures'/(STEM+'.pdf'),17,center_heading=True)
    jwrite(out/'single_page_alignment.json',alignment)
    hist.close()
    print('Rendered one-page ED17: top100/top200 gene overlap and enrichment; historical age fields.',flush=True)

def set_paragraph(doc,id,en,zh):
    found=doc.xpath('//*[@id=$id]',id=id)
    if not found:return
    node=found[0]
    for lang,content in [('en',en),('zh-CN',zh)]:
        ps=node.xpath('./*[@lang=$lang]',lang=lang)
        if not ps:continue
        p=ps[0];refs=[deepcopy(x) for x in p.xpath('.//sup[@data-reference-ids]')]
        for i,r in enumerate(refs):content=content.replace('[[C'+str(i)+']]',html.tostring(r,encoding='unicode',with_tail=False))
        assert '[[C' not in content,(id,lang)
        new=html.fragment_fromstring('<'+p.tag+' lang="'+lang+'" class="'+('en' if lang=='en' else 'zh')+'">'+content+'</'+p.tag+'>')
        p.getparent().replace(p,new)

def revise_document(doc):
    nodes=doc.xpath('//*[@id="ed17"]')
    if nodes:
        section=nodes[0]
        for child in list(section):section.remove(child)
        figure=etree.SubElement(section,'figure')
        etree.SubElement(figure,'h3').text='Extended Data Fig. 17 | '+TITLE
        links=etree.SubElement(figure,'p',attrib={'class':'figure-downloads'})
        for ext in ['pdf','svg','png']:
            a=etree.SubElement(links,'a',href='Extended_Data/figures/'+STEM+'.'+ext);a.text=ext.upper();a.tail=' · '
        for label,href in [('Source workbook','Extended_Data/source_data/Extended_Data_Fig17.xlsx'),('Top100 Metascape results','Source_data/metascape_20261010/top100/metascape_result.xlsx'),('Top200 Metascape results','Source_data/metascape_20261010/top200/metascape_result.xlsx')]:
            a=etree.SubElement(links,'a',href=href);a.text=label;a.tail=' · '
        etree.SubElement(figure,'img',src='Extended_Data/figures/'+STEM+'.png',alt='Extended Data Fig. 17 top100 and top200',loading='lazy')
        cap=etree.SubElement(figure,'figcaption')
        for lang,cls,t in [('en','en',CAPTION.split('\n')[0]),('zh-CN','zh',CAPTION.split('\n')[1])]:etree.SubElement(cap,'p',lang=lang,attrib={'class':cls}).text=t
    set_paragraph(doc,'P049',
        'Method selection also depends on downstream analyses. Table 1 summarizes reported capabilities; these are not all directly benchmarked here. For the forebrain example [[C0]], the eight historical top100 lists remain in Supplementary Table 4. The newly supplied Metascape analyses compare nine configurations, including biVI, at top100 and top200 (<a href="#ed17">Extended Data Fig. 17, top and middle rows</a>; updated source workbook). Gene overlaps and enrichment heatmaps are shown for both list sizes. Lists follow different ranking definitions and several are Radial Glia-focused; biVI ranks norm_spliced_mean differential evidence by Bayes factor, including both directions. These are descriptive comparisons of selected gene sets, not validated rankings of biological importance or method accuracy.',
        'Method selection also depends on the downstream analysis. Table 1 summarizes reported capabilities, which have not all been directly benchmarked. Eight historical top100 lists from the forebrain example [[C0]] remain in Supplementary Table 4. Supplied Metascape results compare nine-configuration top100 and top200 lists, including biVI (<a href="#ed17">Extended Data Fig. 17, upper two rows</a>; updated source workbook), showing gene overlap and enrichment heatmaps at both list sizes. List-ranking definitions differ, and several methods focus on Radial Glia; biVI ranks norm_spliced_mean differential evidence by Bayes factor in both directions. This is a descriptive comparison of selected gene sets, not a validated ranking of biological importance or method accuracy.')
    set_paragraph(doc,'P125','Velocity-associated gene lists and enrichment','Velocity-associated gene lists and enrichment')
    set_paragraph(doc,'P133',
        'Metascape [[C0]] results were imported from the supplied top100 and top200 analysis archives for nine configurations. Gene symbols were mapped to human Entrez IDs using the reported database update of 2026-07-01. The reports specify a whole-genome background, cumulative hypergeometric P values, Benjamini-Hochberg q values, P &lt; 0.01, at least three hits and enrichment factor &gt; 1.5, with term clustering at similarity &gt; 0.3. The supplied overlap diagrams and top-20-cluster heatmaps replace the historical enrichment artwork; mapping counts, complete numerical outputs and archive hashes are supplied. biVI lists were selected from Radial Glia-versus-other-cell norm_spliced_mean results by descending Bayes factor, retaining both LFC signs and source order for ties, without an additional FDR cutoff. Importing these reports does not refit velocity models or reconstruct the previous enrichment database state.',
        'We imported the supplied Metascape [[C0]] analysis packages for nine-configuration top100 and top200 lists. Reports use the database updated on 2026-07-01, map symbols to human Entrez IDs and use a whole-genome background. P values use the cumulative hypergeometric distribution and q values use Benjamini-Hochberg correction; thresholds are P&lt;0.01, at least three hits, enrichment factor&gt;1.5 and term-clustering similarity&gt;0.3. Returned overlap plots and heatmaps of the top 20 enrichment clusters replace the historical enrichment artwork. Mapping counts, complete numerical outputs and archive hashes are supplied. biVI lists use Radial Glia-versus-other-cell norm_spliced_mean differential results, ranked by descending Bayes factor to select 100/200 genes, retaining both LFC signs and source-row order for ties without an additional FDR threshold. Importing these reports neither refits velocity models nor reconstructs the earlier enrichment-database state.')
    for node in doc.xpath('//*[@id="P056"]//*'):
        for attr in ['text','tail']:
            value=getattr(node,attr)
            if value:setattr(node,attr,re.sub(r'17[deg]', '17', re.sub(r',[efh]\b', '', value)))

def update_workbook():
    path=B/'Extended_Data/source_data/Extended_Data_Fig17.xlsx'
    w=openpyxl.load_workbook(path)
    for s in list(w.sheetnames):
        if s.startswith('MS100 ') or s.startswith('MS200 ') or s=='Metascape notes':del w[s]
    for n in RUNS:
        for name,relative in [('Inputs','submitted_gene_lists.csv'),('Mapped genes','mapped_membership.csv'),('List counts','list_counts.csv'),('Heatmap','Enrichment_heatmap/HeatmapSelectedGO.csv'),('All enrichment','Enrichment_GO/GO_AllLists.csv')]:
            s=w.create_sheet(f'MS{n} '+name)
            with (D/f'top{n}'/relative).open() as f:
                for row in csv.reader(f):
                    assert all(len(v)<=32767 for v in row)
                    s.append(row)
            s.freeze_panes='A2'
    s=w.create_sheet('Metascape notes')
    for row in [('field','value'),('Active data','MS100 and MS200 sheets; earlier sheets are preserved historical records'),('Database update','2026-07-01'),('Background','All genes in the genome'),('biVI','Radial Glia vs other cells; norm_spliced_mean; BF descending; both LFC signs'),('Layout','One page: top row top100; middle row top200; bottom row historical age fields; no panel letters; centered headings; shared-function overlap omitted'),('Historical panels','Bottom row: exact 784-gene run unresolved'),('Current caption',CAPTION)]:s.append(row)
    for row in w['Release 1010']:
        if row[0].value=='Current caption':row[1].value=CAPTION
    import tempfile,os
    fd,temporary=tempfile.mkstemp(prefix='.ed17-workbook-',suffix='.xlsx',dir=path.parent);os.close(fd)
    w.save(temporary);Path(temporary).replace(path)

def publish_reports():
    page_count=len(json.loads((B/'Extended_Data/pages.json').read_text()))
    for filename in ['Revision_report.html','Reproduce.html','index.html','Tables_and_source_data_1010.html']:
        path=B/filename;doc=parse(path)
        for old in doc.xpath('//*[@id="ed17-metascape-update"]'):old.getparent().remove(old)
        if filename=='Reproduce.html':
            markup='<section id="ed17-metascape-update"><h2>Extended Data Fig. 17: nine-configuration top100 / top200 Metascape</h2><pre>/usr/bin/python3 Code/ed17_metascape.py --render --publish\n/usr/bin/python3 Code/export_word.py\n/usr/bin/python3 Code/export_browser.py\n/usr/bin/python3 Code/validate_release.py</pre><p>Source_data/metascape_20261010 retains the original analysis reports, numerical tables, mappings, inputs and artwork. The earlier eight-configuration gene table is historical; new data are in source-workbook sheets MS100 / MS200. Age-group fields remain historical artwork.</p></section>'
        else:
            markup='<section id="ed17-metascape-update"><p>Extended Data Fig. 17 is consolidated into one page, excluding shared-function overlap: nine-configuration top100 and top200 Metascape results including biVI. The complete Extended Data set totals 19 pages. <a href="Extended_Data/figures/Extended_Data_Fig17.pdf">Figure 17 PDF</a> · <a href="Extended_Data/source_data/Extended_Data_Fig17.xlsx">Updated source data</a> · <a href="Provenance/ed17_metascape.json">Provenance and mappings</a>. The reports use a whole-genome background and database version 2026-07-01. New enrichment results replace earlier artwork; age-group fields remain historical examples.</p></section>'
        markup=markup.replace('totals 19 pages',f'totals {page_count} pages')
        doc.find('body').insert(2,html.fragment_fromstring(markup))
        for tr in doc.xpath('//tr'):
            cells=tr.xpath('./td')
            if len(cells)==4 and cells[0].text_content()=='extended' and cells[1].text_content()=='17':
                if '_top200' in cells[3].text_content():
                    tr.getparent().remove(tr);continue
                cells[3].text='Extended_Data/figures/'+STEM+'.pdf'
            if tr.xpath('./td') and tr.xpath('./td')[0].text_content().strip()=='17b–c':tr.getparent().remove(tr)
        for node in doc.iter():
            for attr in ['text','tail']:
                value=getattr(node,attr)
                if value:
                    value=re.sub('17 Extended Data figures \\(\\d+ pages\\)',f'17 Extended Data figures ({page_count} pages)',value).replace('17d–e','17e–f').replace('17g–h','17e–f').replace('17e–f','17 bottom row')
                    value=re.sub('Extended Data now includes 17 numbered figures and \\d+ pages',f'Extended Data now includes 17 numbered figures and {page_count} pages',value)
                    value=value.replace('Updated Extended Data figures 6, 7 and 17','Updated Extended Data figures 6 and 7').replace('Restrict aging and historical enrichment results to examples with incomplete provenance','Restrict aging results to examples with incomplete provenance; enrichment has been updated using the supplied Metascape packages')
                    value=value.replace('For missing clone pairs, enrichment backgrounds or aging-run provenance','For missing clone pairs or aging-run provenance')
                    setattr(node,attr,value)
        write(path,doc)

def publish():
    update_workbook()
    rows=[r for r in json.loads((B/'Extended_Data/pages.json').read_text()) if r['figure']!=17]
    rows.append({'figure':17,'v2_figure':19,'part':'combined','filename':STEM,'title':TITLE,'caption':CAPTION})
    jwrite(B/'Extended_Data/pages.json',rows)
    align=json.loads((B/'Provenance/extended_top_alignment.json').read_text())
    align={k:v for k,v in align.items() if v['figure']!=17}
    align['Extended_Data/figures/'+STEM+'.pdf']=json.loads((B/'Reports/ed17_metascape/single_page_alignment.json').read_text())
    jwrite(B/'Provenance/extended_top_alignment.json',align)
    mp=readcsv(B/'Provenance/figure_map.csv');keys=list(mp[0]);mp=[r for r in mp if not(r['kind']=='extended' and r['current_figure']=='17')]
    mp.append(dict(zip(keys,['extended','17','19','Assets/arial/'+STEM+'.pdf','Extended_Data/figures/'+STEM+'.pdf','Single-page top100/top200 gene overlap and enrichment; supplied report geometry, Arial circos labels; historical age fields retained'])))
    csvwrite(B/'Provenance/figure_map.csv',[keys]+[[r[k] for k in keys] for r in mp])
    for name in ['Manuscript_1010_bilingual.html','Manuscript_1010_English_complete.html','Extended_Data_Figures_1010.html']:
        doc=parse(B/name);revise_document(doc)
        if 'English' in name:
            for node in doc.xpath('//*[@lang="zh-CN" and not(self::html)]'):node.getparent().remove(node)
        write(B/name,doc)
    standalone=parse(B/'Manuscript_1010_bilingual.html')
    for im in standalone.xpath('//img'):
        im.set('src','data:image/png;base64,'+base64.b64encode((B/unquote(im.get('src'))).read_bytes()).decode())
    write(B/'Manuscript_1010_bilingual_standalone.html',standalone)
    merged=fitz.open();legends=fitz.open()
    for r in rows:
        with fitz.open(B/'Extended_Data/figures'/(r['filename']+'.pdf')) as doc:merged.insert_pdf(doc)
    for n in range(1,18):
        page=legends.new_page(width=595.276,height=841.89)
        txt='Extended Data Figure '+str(n)+'\n\n'+'\n\n'.join(r['caption'].split('\n')[0] for r in rows if r['figure']==n)
        assert page.insert_textbox(fitz.Rect(36,36,559,805),txt,fontsize=9,fontname='helv',lineheight=1.3)>=0
    savepdf(merged,B/'Extended_Data/Extended_Data_Figures_1010.pdf');savepdf(legends,B/'Extended_Data/Extended_Data_Legends_1010.pdf')
    missing=readcsv(B/'Reports/missing_data.csv');keys=list(missing[0]);missing=[r for r in missing if r['figures']!='17b–c']
    for r in missing:
        if r['figures'] in ['17d–e','17g–h','17e–f']:r['figures']='17 bottom row'
    csvwrite(B/'Reports/missing_data.csv',[keys]+[[r[k] for k in keys] for r in missing])
    publish_reports()
    mapping=readcsv(B/'Provenance/code_result_map.csv');keys=list(mapping[0])
    for row in mapping:
        if row['result']=='Updated Extended Data figures 6, 7 and 17':row['result']='Updated Extended Data figures 6 and 7'
    mapping=[r for r in mapping if r['result']!='Extended Data Fig. 17 Metascape']
    mapping.append(dict(zip(keys,['Extended Data Fig. 17 Metascape','Code/ed17_metascape.py','Source_data/metascape_20261010/','Nine-configuration top100/top200; imported analysis reports and values, retaining historical age-group artwork'])) )
    csvwrite(B/'Provenance/code_result_map.csv',[keys]+[[r[k] for k in keys] for r in mapping])
    review=readcsv(B/'Reports/paragraph_review.csv');keys=list(review[0]);doc=parse(B/'Manuscript_1010_bilingual.html')
    for row in review:
        if row['id'] in ['P049','P056','P125','P133']:
            row['review_action']='Updated for supplied nine-list Metascape top100/top200 results'
            row['english_start']=doc.get_element_by_id(row['id']).xpath('./*[@lang="en"]')[0].text_content()[:120]
    csvwrite(B/'Reports/paragraph_review.csv',[keys]+[[r[k] for k in keys] for r in review])
    for relative in ['completion_status.json','Reports/build_metadata.json']:
        data=json.loads((B/relative).read_text());data.update({'extended_pages':len(rows),'ed17_metascape_imported':True,'ed17_pages':1,'ed17_shared_function_overlap_shown':False,'ed17_configurations':9,'ed17_top_n':[100,200],'ed17_source':'Source_data/metascape_20261010/import_manifest.json','ed17_model_fitting_performed':False,'ed17_updated_at':datetime.now(timezone.utc).isoformat()})
        if relative=='completion_status.json':data['latest_change']='ED17 combined into one page: nine-configuration top100/top200 gene overlap and enrichment; shared-function overlap omitted; historical age-group fields retained.'
        jwrite(B/relative,data)
    print('Published ED17 readers, bilingual text, workbook and merged PDFs.',flush=True)
    from extended_visibility import apply_visibility
    apply_visibility()
    status=json.loads((B/'completion_status.json').read_text())
    status['latest_change']='ED17 panel letters removed; Top100/Top200 headings shortened and headings centered.'
    jwrite(B/'completion_status.json',status)
    # Retain the release-wide A4 legend delivery when it is active.
    from build_a4_figures import refresh_published
    refresh_published()

if __name__=='__main__':
    if '--import' in sys.argv:import_data()
    if '--render' in sys.argv:render()
    if '--publish' in sys.argv:publish()
